"""Сборка принятых карьерных карт: проверки, арифметика и оформление."""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import math
import shutil
import subprocess
import tempfile
import textwrap
from collections import Counter
from copy import copy
from html import unescape
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

TRACKS = {"project": "Проектный трек", "commercial": "Коммерческий трек"}
TRACK_NAMES = {**TRACKS, "exclude": "Исключена"}
BANDS = {"core": "Общий навык", "specialty": "Специализация", "industry": "Отраслевое требование"}
STATUSES = {"required": "Требование", "task": "Задача", "nice": "Будет плюсом", "unclear": "Неоднозначно"}
COLORS = {"project": "18364F", "commercial": "126D70"}
INK, WHITE, STRIPE, WARM = "213547", "FFFFFF", "F3F6F8", "FFF1D8"
DEGREE_IDS = {
    "project_033", "project_038", "project_195",
    "commercial_023", "commercial_204", "commercial_265",
}
HEADER, FIRST, FLAG, WEIGHT, VALID, ID_COL = 9, 10, 10, 11, 12, 13
HEADERS = [
    "Навык", "Что оценить", "Примеры из вакансий", "Вакансий", "Доля",
    "Самооценка", "Вклад", "Комментарий", "Действия",
    "В общей оценке", "Вес", "Оценка заполнена", "ID навыка",
]
GUIDE = [
    "Частота учитывает задачи, требования и общие пожелания. Для отраслей учтены только обязательные условия.",
    "Шкала 1-5. 1 - нет опыта; 2 - нужна помощь; 3 - работаю самостоятельно; 4 - решаю сложные задачи; 5 - могу обучать.",
    "Вклад зависит от оценки и частоты навыка. Оценка 1 даёт нулевой вклад, 3 - половину веса, 5 - полный вес.",
    "В итог входят 20 самых частых общих навыков. Специализации и отраслевые требования оцените отдельно.",
    "Заполните светло-жёлтые ячейки. В комментарии приведите свой пример, в действиях запишите следующий шаг. "
    "Строки с цветной полосой у оценки входят в итог. Для общей оценки нужны все 20 оценок.",
]
EVIDENCE_HEADERS = [
    "Трек", "Навык", "Статус", "Цитата", "Компания", "Вакансия", "Ссылка",
    "В частоте", "Раздел источника", "ID вакансии", "ID навыка",
]
LEDGER_HEADERS = [
    "Компания", "Вакансия", "Трек", "Активна", "Ссылка",
    "Исходная ссылка", "Формулировка роли", "ID вакансии",
]


def require(condition, message, detail=None):
    if not condition:
        raise ValueError(f"{message}: {detail}" if detail is not None else message)


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def eligible(req):
    return req["status"] in ("required", "task", "nice") and (
        req["category"] != "industry" or req["status"] == "required"
    )


def evidence_key(track, req):
    return json.dumps([track, req["category"], req["concept"]], ensure_ascii=False, separators=(",", ":"))


def signature(item):
    return str(item["id"]), item["quote"], item["scope"]


def quoted_text(row, scope):
    if scope == "description":
        return row.get("description") or ""
    require(scope.startswith("native:"), "Неизвестный раздел цитаты", scope)
    return (row.get("native_sections") or {}).get(scope[7:]) or ""


def load_data(original, maps):
    manifest = read_json(original / "accepted.json")
    require(manifest["meta"]["total"] == 342, "Изменился размер среза")
    require(manifest["expected_tracks"] == {"project": 183, "commercial": 118, "exclude": 41},
            "Изменились принятые треки")
    for filename, expected in manifest["data_blobs"].items():
        payload = (original / filename).read_bytes()
        actual = hashlib.sha1(f"blob {len(payload)}\0".encode() + payload).hexdigest()
        require(actual == expected, "Изменился принятый исходник", filename)
    rows, decisions = [], []
    for filename in manifest["source_files"]:
        rows.extend(read_json(original / filename)["rows"])
    for filename in manifest["analysis_files"]:
        decisions.extend(read_json(original / filename))
    source_updates = read_json(original / manifest["source_overrides_file"])
    updates = {str(item["id"]): item["fields"] for item in source_updates}
    require(len(updates) == len(source_updates), "Дубли исходных исправлений")
    by_id = {str(row["id"]): row for row in rows}
    require(set(updates) <= set(by_id), "Исправления неизвестных вакансий")
    # Сначала применяем исправления источников, затем проверяем цитаты и формулировки роли.
    for row in rows:
        row.update(updates.get(str(row["id"]), {}))
    review = read_json(original / manifest["review_file"])
    require(review["status"] == "ACCEPT", "Разбор не принят")
    replacements = {str(item["id"]): item for item in review["replacements"]}
    require(len(replacements) == len(review["replacements"]), "Дубли исправлений разбора")
    decisions = [replacements.get(str(item["id"]), item) for item in decisions]
    by_decision = {str(item["id"]): item for item in decisions}
    require(len(rows) == len(by_id) == len(decisions) == len(by_decision) == 342,
            "Повтор или потеря вакансии")
    require(set(by_id) == set(by_decision), "Не совпадают вакансии и решения")
    counts = Counter(item["track"] for item in decisions)
    require(dict(counts) == manifest["expected_tracks"], "Изменилась классификация", dict(counts))
    require(set(replacements) <= set(by_decision), "Исправления неизвестных решений")

    evidence = []
    for decision in decisions:
        row = by_id[str(decision["id"])]
        bodies = [row.get("description") or "", *(row.get("native_sections") or {}).values()]
        require(decision["role_evidence"] and any(
            decision["role_evidence"] in body for body in bodies if isinstance(body, str)
        ), "Формулировка роли отсутствует в источнике", decision["id"])
        if decision["track"] == "exclude":
            require(not decision["requirements"], "У исключённой роли есть требования", decision["id"])
            continue
        for index, req in enumerate(decision["requirements"]):
            require(req["status"] in STATUSES, "Неизвестный статус", (decision["id"], index))
            require(req["quote"] and req["quote"] in quoted_text(row, req["scope"]),
                    "Цитата не совпадает с исходником", (decision["id"], index))
            evidence.append({**req, "id": str(decision["id"]), "index": index,
                             "track": decision["track"], "key": evidence_key(decision["track"], req)})
    require(len(evidence) == manifest["expected_requirements"], "Изменилось число исходных требований")

    old_groups = read_json(original / "taxonomy.json")["groups"]
    old_by_id = {group["id"]: group for group in old_groups}
    require(len(old_by_id) == len(old_groups) == manifest["accepted_groups"] == 671,
            "Повтор или потеря исходной группы")
    old_mapping = {}
    for old in old_groups:
        for key in old["member_keys"]:
            require(key not in old_mapping, "Требование повторяется в исходной карте", key)
            require(json.loads(key)[0] == old["track"], "Требование другого трека", old["id"])
            old_mapping[key] = old["id"]
    require(set(old_mapping) == {item["key"] for item in evidence if eligible(item)},
            "Исходная карта не покрывает принятые требования")
    for item in evidence:
        item["old_id"] = old_mapping.get(item["key"])
    old_evidence = {old_id: [] for old_id in old_by_id}
    for item in evidence:
        if item["old_id"]:
            old_evidence[item["old_id"]].append(item)

    new_groups, old_to_new, excluded = [], {}, set()
    for track in TRACKS:
        catalogue = read_json(maps / f"catalogue_{track}.json")
        compact = read_json(maps / f"compact_{track}.json")
        require(catalogue["track"] == compact["track"] == track, "Карта другого трека", track)
        require(catalogue["total_vacancies"] == counts[track], "Неверная база каталога", track)
        removed = set(catalogue["excluded_degree_ids"])
        require(len(removed) == len(catalogue["excluded_degree_ids"]) == 3,
                "Повтор или потеря исключённой группы", track)
        excluded.update(removed)
        cat_groups = catalogue["groups"]
        cat_by_id = {group["id"]: group for group in cat_groups}
        expected_ids = {old["id"] for old in old_groups if old["track"] == track} - removed
        require(set(cat_by_id) == expected_ids and len(cat_by_id) == len(cat_groups),
                "Каталог не совпадает с исходной картой", track)
        for old_id, cat in cat_by_id.items():
            old = old_by_id[old_id]
            require(cat["category"] == old["category"], "Изменился тип исходной группы", old_id)
            require(set(cat["concepts"]) == {json.loads(key)[2] for key in old["member_keys"]},
                    "Изменился состав исходной группы", old_id)
            known = {signature(item) for item in old_evidence[old_id]}
            require(cat["count"] == len({item["id"] for item in old_evidence[old_id] if eligible(item)}),
                    "Частота каталога не совпадает с исходными фактами", old_id)
            for example in cat["evidence"]:
                require(signature(example) in known, "Пример каталога отсутствует в исходной группе", old_id)
                require(example["company"] == by_id[str(example["id"])]["company"],
                        "Компания каталога не совпадает с исходником", old_id)
        require(sum(group["band"] == "core" for group in compact["groups"]) >= 20,
                "Нужны хотя бы 20 общих навыков", track)
        covered = []
        for group in compact["groups"]:
            require(group["id"] and group["label"].strip() and group["description"].strip(),
                    "Пустой навык", group.get("id"))
            require(group["band"] in BANDS, "Неизвестный раздел", group["id"])
            source_ids = group["source_group_ids"]
            require(source_ids and len(source_ids) == len(set(source_ids)),
                    "Пустой или повторный состав нового навыка", group["id"])
            require(set(source_ids) <= set(cat_by_id), "Чужая или исключённая группа", group["id"])
            is_industry = [old_by_id[old_id]["category"] == "industry" for old_id in source_ids]
            require(all(is_industry) if group["band"] == "industry" else not any(is_industry),
                    "Отраслевое условие смешано с практикой", group["id"])
            covered.extend(source_ids)
            local_evidence = [item for old_id in source_ids for item in old_evidence[old_id]]
            counted = [item for item in local_evidence if eligible(item)]
            ids = {item["id"] for item in counted}
            require(ids, "Навык без учитываемых вакансий", group["id"])
            known = {signature(item) for item in counted}
            examples = group["examples"]
            require(1 <= len(examples) <= 3 and len({ex["quote"] for ex in examples}) == len(examples),
                    "Нужны от одного до трёх разных точных примеров", group["id"])
            require(len(examples) == min(3, len({item["quote"] for item in counted})),
                    "Не хватает доступных точных примеров", group["id"])
            for example in examples:
                require(signature(example) in known, "Пример не принадлежит навыку", (group["id"], example))
                require(example["company"] == by_id[str(example["id"])]["company"],
                        "Компания примера не совпадает с источником", group["id"])
            for old_id in source_ids:
                require(old_id not in old_to_new, "Исходная группа учтена повторно", old_id)
                old_to_new[old_id] = group["id"]
            new_groups.append({**group, "track": track, "count": len(ids), "vacancy_ids": ids})
        require(set(covered) == expected_ids and len(covered) == len(expected_ids),
                "Карта не покрывает каталог ровно один раз", track)
    require(excluded == DEGREE_IDS, "Изменился список исключённых групп", sorted(excluded))
    require(len(old_to_new) == 665 and set(old_to_new) == set(old_by_id) - DEGREE_IDS,
            "Не сохранены все 665 исходных групп")
    new_by_id = {group["id"]: group for group in new_groups}
    require(len(new_by_id) == len(new_groups), "Повтор ID новых навыков")
    exported = []
    for item in evidence:
        if item["old_id"] in excluded:
            continue
        new_id = old_to_new.get(item["old_id"])
        exported.append({**item, "new_id": new_id,
                         "counted": bool(new_id and eligible(item))})
    ordered = {}
    for track in TRACKS:
        ordered[track] = sorted(
            [group for group in new_groups if group["track"] == track],
            key=lambda group: (list(BANDS).index(group["band"]), -group["count"], group["id"]),
        )
        core = [group for group in ordered[track] if group["band"] == "core"]
        scored_ids = {group["id"] for group in core[:20]}
        denominator = sum(group["count"] for group in core[:20])
        for group in ordered[track]:
            group["scored"] = group["id"] in scored_ids
            group["weight"] = group["count"] / denominator if group["scored"] else 0
    return {"manifest": manifest, "rows": rows, "decisions": by_decision, "by_id": by_id,
            "counts": counts, "groups": ordered, "by_group": new_by_id, "evidence": exported,
            "removed_evidence": len(evidence) - len(exported), "source_groups": len(old_to_new)}


def put(ws, row, col, value, fill=None, bold=False, color=INK, size=11):
    cell = ws.cell(row, col, value)
    # Цитаты и названия сохраняем буквально. Знак "=" в начале не превращает текст в формулу.
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Calibri", size=size, color=color, bold=bold)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def formula(ws, row, col, value, number_format=None):
    cell = put(ws, row, col, value)
    cell.data_type = "f"
    require(value.isascii(), "Формула должна использовать ASCII", value)
    if number_format:
        cell.number_format = number_format
    return cell


def merged(ws, region, value, fill=None, bold=False, size=11, color=INK):
    ws.merge_cells(region)
    cell = ws[region.split(":")[0]]
    put(ws, cell.row, cell.column, value, fill, bold, color, size)


def common_setup(ws, color, title, last=9):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = color
    merged(ws, f"A1:{ws.cell(1, last).column_letter}1", title, color, True, 20, WHITE)
    ws.row_dimensions[1].height = 35
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.20, right=0.20, top=0.20, bottom=0.20, header=0, footer=0)


def row_height(text, width, line=14):
    return sum(max(1, len(textwrap.wrap(part, width=width))) for part in text.split("\n")) * line + 14


def row_formulas(ws, row):
    formula(ws, row, VALID,
            f'=IFERROR(IF(AND(ISNUMBER(F{row}),F{row}>=1,F{row}<=5,MOD(F{row},1)=0),1,0),0)')
    contribution = formula(ws, row, 7,
                           f'=IF(J{row}=0,"",IF(L{row}=1,(F{row}-1)/4*K{row},""))', "0.0%")
    contribution.fill = PatternFill("solid", fgColor=STRIPE if row % 2 == 0 else WHITE)


def track_caption(data, track):
    return (f'Срез 14 октября 2026 · {data["counts"][track]} вакансий в сохранённом пуле, включая архивные. '
            + GUIDE[0])


def build_workbook(data, output):
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    wb.properties.title = "Карьерный шкаф"
    wb.properties.subject = "Проектный и коммерческий треки по срезу 14 октября 2026"
    wb.properties.creator = "ProdRadar"
    for track, title in TRACKS.items():
        ws = wb.create_sheet(title)
        common_setup(ws, COLORS[track], title)
        last = FIRST + len(data["groups"][track]) - 1
        merged(ws, "A2:I2", track_caption(data, track))
        ws.row_dimensions[2].height = 34
        for row, message in [(3, GUIDE[1]), (4, GUIDE[2]), (5, GUIDE[3]), (7, GUIDE[4])]:
            merged(ws, f"A{row}:I{row}", message)
            ws.row_dimensions[row].height = 31
        merged(ws, "A6:B6", "Оценено навыков", STRIPE, True, 12)
        formula(ws, 6, 3, f"=SUMPRODUCT(J{FIRST}:J{last},L{FIRST}:L{last})", "0")
        merged(ws, "D6:E6", "из 20", STRIPE, True, 12)
        merged(ws, "F6:G6", "Соответствие по самооценке", STRIPE, True, 12)
        ws.merge_cells("H6:I6")
        formula(ws, 6, 8, f'=IF(C6=20,SUM(G{FIRST}:G{last}),"")', "0.0%")
        for col in (3, 8):
            ws.cell(6, col).font = Font(name="Calibri", size=18, bold=True, color=COLORS[track])
            ws.cell(6, col).fill = PatternFill("solid", fgColor="E8F0F3")
        ws.row_dimensions[6].height = 43
        merged(ws, "A8:I8",
               "Сначала общие навыки, затем специализации и отраслевые требования. "
               "Первые 20 общих навыков входят в общую оценку; остальные строки помогают спланировать развитие.",
               bold=True)
        ws.row_dimensions[8].height = 30
        for col, value in enumerate(HEADERS, 1):
            put(ws, HEADER, col, value, COLORS[track], True, WHITE)
        ws.row_dimensions[HEADER].height = 34
        for row, group in enumerate(data["groups"][track], FIRST):
            quotes = "\n\n".join(f'{ex["company"]} - {unescape(ex["quote"])}' for ex in group["examples"])
            values = [group["label"], group["description"], quotes, group["count"],
                      group["count"] / data["counts"][track], None, None, None, None,
                      int(group["scored"]), group["weight"], None, group["id"]]
            for col, value in enumerate(values, 1):
                put(ws, row, col, value,
                    WARM if col in (6, 8, 9) else (STRIPE if row % 2 == 0 else WHITE),
                    bold=col == 1)
            ws.cell(row, 4).number_format = "0"
            ws.cell(row, 5).number_format = "0.0%"
            ws.cell(row, WEIGHT).number_format = "0.000000%"
            row_formulas(ws, row)
            ws.cell(row, 1).comment = Comment(
                BANDS[group["band"]] + (". Входит в общую оценку." if group["scored"] else ". Для плана развития."),
                "ProdRadar",
            )
            if group["scored"]:
                ws.cell(row, 1).fill = PatternFill("solid", fgColor="E6EEF4" if track == "project" else "E4F0EF")
                ws.cell(row, 1).font = Font(name="Calibri", size=11, bold=True, color=COLORS[track])
                ws.cell(row, 6).border = Border(left=Side(style="medium", color=COLORS[track]))
            ws.row_dimensions[row].height = min(409, max(
                105, row_height(group["label"], 26), row_height(group["description"], 40),
                row_height(quotes, 58), 105,
            ))
        table = Table(displayName=f"Career_{track}", ref=f"A{HEADER}:M{last}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2" if track == "project" else "TableStyleMedium9",
            showFirstColumn=False, showLastColumn=False, showRowStripes=False, showColumnStripes=False,
        )
        ws.add_table(table)
        ws.auto_filter.ref = table.ref
        dv = DataValidation(type="whole", operator="between", formula1=1, formula2=5, allow_blank=True)
        dv.error, dv.errorTitle = "Введите целое число от 1 до 5.", "Самооценка"
        dv.prompt, dv.promptTitle = "Оцените свой опыт по шкале вверху листа.", "Оценка опыта"
        dv.showErrorMessage = dv.showInputMessage = True
        dv.errorStyle = "stop"
        ws.add_data_validation(dv)
        dv.add(f"F{FIRST}:F{last}")
        for col, width in {"A": 27, "B": 42, "C": 60, "D": 10, "E": 10,
                           "F": 9, "G": 10, "H": 28, "I": 28}.items():
            ws.column_dimensions[col].width = width
        for col in ("J", "K", "L", "M"):
            ws.column_dimensions[col].hidden = True
        ws.freeze_panes = f"C{FIRST}"
        ws.print_title_rows = f"1:{HEADER}"
        ws.print_area = f"A1:I{last}"

    ws = wb.create_sheet("Цитаты")
    common_setup(ws, "647787", "Точные формулировки из вакансий", len(EVIDENCE_HEADERS))
    merged(ws, "A2:K2", "Все сохранённые цитаты и статусы. Пустой навык означает, что требование не участвует в карте.")
    ws.row_dimensions[2].height = 30
    for col, value in enumerate(EVIDENCE_HEADERS, 1):
        put(ws, 4, col, value, COLORS["project"], True, WHITE)
    for row, item in enumerate(data["evidence"], 5):
        job = data["by_id"][item["id"]]
        group = data["by_group"].get(item["new_id"])
        values = [TRACKS[item["track"]], group["label"] if group else None,
                  STATUSES[item["status"]], item["quote"], job["company"], job["title"],
                  job["url"], "Да" if item["counted"] else "Нет", item["scope"], item["id"], item["new_id"]]
        for col, value in enumerate(values, 1):
            put(ws, row, col, value, STRIPE if row % 2 == 0 else WHITE)
        ws.cell(row, 7).hyperlink = job["url"]
        ws.row_dimensions[row].height = min(409, max(44, row_height(item["quote"], 86)))
    for col, width in {"A": 22, "B": 35, "C": 18, "D": 90, "E": 22, "F": 50,
                       "G": 35, "H": 12, "I": 25, "J": 42, "K": 23}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes, ws.auto_filter.ref = "D5", f"A4:K{ws.max_row}"

    ws = wb.create_sheet("Вакансии")
    common_setup(ws, "83919A", "Исходные вакансии", len(LEDGER_HEADERS))
    merged(ws, "A2:H2", "342 вакансии, включая 41 исключённую. Названия и формулировки роли сохранены из источников.")
    ws.row_dimensions[2].height = 30
    for col, value in enumerate(LEDGER_HEADERS, 1):
        put(ws, 4, col, value, COLORS["project"], True, WHITE)
    for row, job in enumerate(data["rows"], 5):
        decision = data["decisions"][str(job["id"])]
        values = [job["company"], job["title"], TRACK_NAMES[decision["track"]],
                  "Да" if job["is_active"] else "Нет", job["url"],
                  job.get("original_url", job["url"]), decision["role_evidence"], str(job["id"])]
        for col, value in enumerate(values, 1):
            put(ws, row, col, value, STRIPE if row % 2 == 0 else WHITE)
        for col in (5, 6):
            ws.cell(row, col).hyperlink = values[col - 1]
        ws.row_dimensions[row].height = min(409, max(60, row_height(values[6], 80)))
    for col, width in {"A": 22, "B": 57, "C": 23, "D": 10,
                       "E": 40, "F": 40, "G": 85, "H": 43}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes, ws.auto_filter.ref = "C5", f"A4:H{ws.max_row}"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    verify_workbook(output, data)


def verify_workbook(path, data, cached=False):
    wb = load_workbook(path, data_only=False)
    require(wb.sheetnames == [*TRACKS.values(), "Цитаты", "Вакансии"], "Неверные листы Excel")
    require(wb["Вакансии"].max_row - 4 == 342, "Потеря вакансий в Excel")
    require(wb["Цитаты"].max_row - 4 == len(data["evidence"]), "Потеря цитат в Excel")
    for track, title in TRACKS.items():
        ws = wb[title]
        groups = data["groups"][track]
        require(ws.max_row - HEADER == len(groups), "Потеря навыков в Excel", track)
        scored_ids, expected_ids = set(), {group["id"] for group in groups}
        for row in range(FIRST, ws.max_row + 1):
            group = data["by_group"][ws.cell(row, ID_COL).value]
            expected_quotes = "\n\n".join(f'{ex["company"]} - {unescape(ex["quote"])}' for ex in group["examples"])
            require(ws.cell(row, 3).value == expected_quotes, "Изменился пример в Excel", group["id"])
            require(ws.cell(row, 4).value == len(group["vacancy_ids"]), "Неверный union count", group["id"])
            require(math.isclose(ws.cell(row, 5).value, group["count"] / data["counts"][track]),
                    "Неверная доля", group["id"])
            require(all(ws.cell(row, col).value is None for col in (6, 8, 9)), "Предзаполнены поля ввода")
            if ws.cell(row, FLAG).value == 1:
                scored_ids.add(group["id"])
            require(ws.cell(row, FLAG).value == int(group["scored"]), "Изменился набор самооценки")
            require(math.isclose(ws.cell(row, WEIGHT).value, group["weight"]), "Изменился вес навыка")
            for col in (7, VALID):
                require(ws.cell(row, col).data_type == "f" and ws.cell(row, col).value.isascii(),
                        "Формула не сохранена", (title, row, col))
        require(len(scored_ids) == 20 and scored_ids == {group["id"] for group in groups if group["scored"]},
                "Неверный набор 20 навыков", track)
        require(math.isclose(sum(ws.cell(row, WEIGHT).value for row in range(FIRST, ws.max_row + 1)),
                             1.0, abs_tol=1e-12), "Сумма весов отличается от единицы", track)
        require({ws.cell(row, ID_COL).value for row in range(FIRST, ws.max_row + 1)} == expected_ids,
                "ID навыков не совпадают")
    for row, item in enumerate(data["evidence"], 5):
        require(wb["Цитаты"].cell(row, 4).value == item["quote"], "Изменилась цитата в Excel", row)
    wb.close()
    if cached:
        checked = load_workbook(path, data_only=True)
        for title in TRACKS.values():
            ws = checked[title]
            require(ws["C6"].value == 0 and ws["H6"].value in (None, ""), "Неверное начальное состояние")
            for row in range(FIRST, ws.max_row + 1):
                require(ws.cell(row, VALID).value == 0, "Нет пересчитанного значения формулы")
                require(ws.cell(row, 7).value in (None, ""), "Неверный пустой вклад")
        checked.close()


def run_calc(source, destination, pdf=False):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    require(executable, "LibreOffice Calc не установлен")
    destination.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="career-calc-profile-") as profile:
        result = subprocess.run(
            [executable, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless",
             "--convert-to", "pdf:calc_pdf_Export" if pdf else "xlsx:Calc MS Excel 2007 XML",
             "--outdir", str(destination), str(source)],
            check=False, capture_output=True, text=True, timeout=180,
        )
    target = destination / f'{source.stem}.{"pdf" if pdf else "xlsx"}'
    require(result.returncode == 0 and target.is_file() and target.stat().st_size > 0,
            "LibreOffice не создал файл", result.stdout + result.stderr)
    return target


def recalculate(path, data):
    with tempfile.TemporaryDirectory(prefix="career-recalculated-") as temporary:
        recalculated = run_calc(path.resolve(), Path(temporary))
        shutil.copy2(recalculated, path)
    # После Calc не сохраняем файл через openpyxl, чтобы не потерять рассчитанные значения.
    verify_workbook(path, data, cached=True)


def formula_qa(path):
    wb = load_workbook(path)
    checks = []
    scenarios = ("blank", "all1", "all3", "all5", "partial", "unequal", "unscored", "sorted",
                 "invalid_decimal", "invalid_zero", "invalid_six", "invalid_text")
    for track, title in TRACKS.items():
        source = wb[title]
        for scenario in scenarios:
            ws = wb.copy_worksheet(source)
            ws.title = f"{track}_{scenario}"
            scored = [row for row in range(FIRST, ws.max_row + 1) if ws.cell(row, FLAG).value == 1]
            unscored = [row for row in range(FIRST, ws.max_row + 1) if ws.cell(row, FLAG).value == 0]
            require(len(scored) == 20, "QA: неверный набор общих навыков")
            if scenario == "sorted":
                records = [[copy(ws.cell(row, col)) for col in range(1, ID_COL + 1)]
                           for row in range(FIRST, ws.max_row + 1)]
                for row, cells in enumerate(reversed(records), FIRST):
                    for col, old_cell in enumerate(cells, 1):
                        new_cell = ws.cell(row, col)
                        new_cell.value = old_cell.value
                        if old_cell.data_type == "f":
                            new_cell.value = Translator(old_cell.value, origin=old_cell.coordinate).translate_formula(new_cell.coordinate)
                        new_cell._style = copy(old_cell._style)
                scored = [row for row in range(FIRST, ws.max_row + 1) if ws.cell(row, FLAG).value == 1]
            require(math.isclose(sum(ws.cell(row, WEIGHT).value for row in scored),
                                 1.0, abs_tol=1e-12), "QA: сумма весов отличается от единицы", ws.title)
            count_total = sum(ws.cell(row, 4).value for row in scored)
            reference_weights = {row: ws.cell(row, 4).value / count_total for row in scored}
            ratings, invalid_rows = {}, set()
            if scenario in ("all1", "all3", "all5"):
                ratings = {row: int(scenario[-1]) for row in scored}
            elif scenario == "partial":
                ratings = {row: 5 for row in scored[:-1]}
            elif scenario in ("unequal", "sorted"):
                require(len({ws.cell(row, WEIGHT).value for row in scored}) > 1,
                        "QA: проверка разных весов требует разных частот", title)
                ratings = {row: 1 + index % 5 for index, row in enumerate(scored)}
            elif scenario == "unscored":
                require(unscored, "QA: нет навыка вне общей оценки", title)
                ratings = {row: 3 for row in scored}
                ratings[unscored[0]] = 5
            elif scenario.startswith("invalid_"):
                ratings = {row: 3 for row in scored}
                invalid_values = {"invalid_decimal": 2.5, "invalid_zero": 0,
                                  "invalid_six": 6, "invalid_text": "опыт"}
                ratings[scored[0]] = invalid_values[scenario]
                invalid_rows.add(scored[0])
            for row, rating in ratings.items():
                ws.cell(row, 6, rating)
            progress = sum(row in ratings and row not in invalid_rows for row in scored)
            canonical = {"all1": 0.0, "all3": 0.5, "all5": 1.0, "unscored": 0.5}
            expected = canonical.get(scenario)
            if progress == 20 and scenario not in canonical:
                expected = sum((ratings[row] - 1) / 4 * reference_weights[row] for row in scored)
            checks.append({"sheet": ws.title, "progress": progress, "expected": expected,
                           "ratings": ratings, "invalid_rows": invalid_rows, "scored": scored,
                           "weights": reference_weights})
    for title in [*TRACKS.values(), "Цитаты", "Вакансии"]:
        wb.remove(wb[title])
    with tempfile.TemporaryDirectory(prefix="career-formula-qa-") as temporary:
        root = Path(temporary)
        input_path = root / "formula_tests.xlsx"
        wb.save(input_path)
        wb.close()
        output = run_calc(input_path, root / "calc")
        cached = load_workbook(output, data_only=True)
        for check in checks:
            ws = cached[check["sheet"]]
            require(ws["C6"].value == check["progress"], "QA: неверный прогресс", check["sheet"])
            actual = ws["H6"].value
            if check["expected"] is None:
                require(actual in (None, ""), "QA: общая оценка появилась раньше времени", check["sheet"])
            else:
                require(isinstance(actual, (int, float)) and math.isclose(actual, check["expected"], abs_tol=1e-9),
                        "QA: неверная взвешенная оценка", (check["sheet"], actual, check["expected"]))
            for row in range(FIRST, ws.max_row + 1):
                valid = row in check["ratings"] and row not in check["invalid_rows"]
                require(ws.cell(row, VALID).value == int(valid),
                        "QA: неверно распознана допустимая оценка", (check["sheet"], row))
                if row in check["scored"] and valid:
                    expected = (check["ratings"][row] - 1) / 4 * check["weights"][row]
                    require(isinstance(ws.cell(row, 7).value, (int, float))
                            and math.isclose(ws.cell(row, 7).value, expected, abs_tol=1e-9),
                            "QA: неверный вклад строки", (check["sheet"], row))
                else:
                    require(ws.cell(row, 7).value in (None, ""), "QA: учтён навык без веса или оценки")
        cached.close()
    return [{"scenario": check["sheet"], "progress": check["progress"], "score": check["expected"]}
            for check in checks]


def preview(path, output):
    require(shutil.which("pdftoppm"), "pdftoppm не установлен")
    previews = []
    with tempfile.TemporaryDirectory(prefix="career-preview-") as temporary:
        root = Path(temporary)
        for track, title in TRACKS.items():
            wb = load_workbook(path)
            for ws in list(wb.worksheets):
                if ws.title != title:
                    wb.remove(ws)
            ws = wb[title]
            ws.print_area = "A1:I12"
            ws.print_title_rows = None
            ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            copy_path = root / f"preview_{track}.xlsx"
            wb.save(copy_path)
            wb.close()
            pdf = run_calc(copy_path, root / track, pdf=True)
            pdf_target = output / f"preview_{track}.pdf"
            shutil.copy2(pdf, pdf_target)
            png_prefix = output / f"preview_{track}"
            result = subprocess.run(
                ["pdftoppm", "-f", "1", "-singlefile", "-scale-to-x", "1600",
                 "-scale-to-y", "-1", "-png", str(pdf), str(png_prefix)],
                check=False, capture_output=True, text=True, timeout=60,
            )
            png = png_prefix.with_suffix(".png")
            require(result.returncode == 0 and png.is_file(), "Не создан PNG", result.stderr)
            previews.append({"track": track, "file": png.name,
                             "image": "data:image/png;base64," + base64.b64encode(png.read_bytes()).decode("ascii")})
    (output / "preview.json").write_text(json.dumps(previews, ensure_ascii=False, separators=(",", ":")),
                                         encoding="utf-8")
    return [item["file"] for item in previews]


def own_text(data):
    texts = [*TRACK_NAMES.values(), *BANDS.values(), *STATUSES.values(), *HEADERS, *GUIDE,
             *EVIDENCE_HEADERS, *LEDGER_HEADERS,
             "Оценено навыков", "из 20", "Соответствие по самооценке",
             "Введите целое число от 1 до 5.", "Оцените свой опыт по шкале вверху листа.",
             "Оценка опыта", "Точные формулировки из вакансий", "Исходные вакансии",
             "Входит в общую оценку.", "Для плана развития.",
             "Все сохранённые цитаты и статусы. Пустой навык означает, что требование не участвует в карте.",
             "342 вакансии, включая 41 исключённую. Названия и формулировки роли сохранены из источников.",
             "Сначала общие навыки, затем специализации и отраслевые требования. "
             "Первые 20 общих навыков входят в общую оценку; остальные строки помогают спланировать развитие."]
    texts.extend(track_caption(data, track) for track in TRACKS)
    texts.extend(("Карьерный шкаф", "Проектный и коммерческий треки по срезу 14 октября 2026", "Да", "Нет"))
    texts.extend(band + suffix for band in BANDS.values()
                 for suffix in (". Входит в общую оценку.", ". Для плана развития."))
    for groups in data["groups"].values():
        for group in groups:
            texts.extend((group["label"], group["description"]))
    return "\n\n".join(dict.fromkeys(texts)) + "\n"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--original", type=Path, default=Path("original/.github/career_closet"))
    parser.add_argument("--maps", type=Path, default=Path(".github/career_closet_v2"))
    parser.add_argument("--output", type=Path, default=Path("output/career_closet_v2"))
    parser.add_argument("--recalculate", action="store_true")
    parser.add_argument("--qa", action="store_true")
    parser.add_argument("--preview", action="store_true")
    args = parser.parse_args()
    data = load_data(args.original, args.maps)
    args.output.mkdir(parents=True, exist_ok=True)
    path = args.output / "Карьерный_шкаф.xlsx"
    build_workbook(data, path)
    (args.output / "own_text.txt").write_text(own_text(data), encoding="utf-8")
    tests = formula_qa(path) if args.qa else []
    if args.recalculate:
        recalculate(path, data)
    previews = preview(path, args.output) if args.preview else []
    report = {"vacancies": len(data["rows"]), "tracks": dict(data["counts"]),
              "source_groups": data["source_groups"],
              "groups": {track: len(groups) for track, groups in data["groups"].items()},
              "quotes": len(data["evidence"]), "removed_degree_quotes": data["removed_evidence"],
              "formula_tests": tests, "recalculated": args.recalculate, "previews": previews}
    (args.output / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print("CAREER_V2=" + json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
