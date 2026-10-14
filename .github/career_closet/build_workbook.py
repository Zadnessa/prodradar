"""Сборка Excel из принятого ручного разбора; смысловые решения здесь не принимаются."""
import json
import hashlib
from html import unescape
from pathlib import Path
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(".github/career_closet")
manifest = json.loads((ROOT / "accepted.json").read_text(encoding="utf-8"))
for filename, expected_blob in manifest["data_blobs"].items():
    payload = (ROOT / filename).read_bytes()
    actual_blob = hashlib.sha1(f"blob {len(payload)}\0".encode() + payload).hexdigest()
    assert actual_blob == expected_blob, ("Исходник изменился после принятия", filename)
taxonomy = json.loads((ROOT / "taxonomy.json").read_text(encoding="utf-8"))
source = {"meta": manifest["meta"]}
rows = []
decisions = []
for filename in manifest["source_files"]:
    rows.extend(json.loads((ROOT / filename).read_text(encoding="utf-8"))["rows"])
for filename in manifest["analysis_files"]:
    decisions.extend(json.loads((ROOT / filename).read_text(encoding="utf-8")))
review = json.loads((ROOT / manifest["review_file"]).read_text(encoding="utf-8"))
assert review["status"] == "ACCEPT"
replacements = {str(d["id"]): d for d in review["replacements"]}
assert len(replacements) == len(review["replacements"])
decisions = [replacements.get(str(d["id"]), d) for d in decisions]
source_overrides = json.loads((ROOT / manifest["source_overrides_file"]).read_text(encoding="utf-8"))
row_updates = {str(r["id"]): r["fields"] for r in source_overrides}
assert len(row_updates) == len(source_overrides)
for row in rows:
    row.update(row_updates.get(str(row["id"]), {}))
groups = taxonomy["groups"]
by_id = {str(r["id"]): r for r in rows}
by_decision = {str(r["id"]): r for r in decisions}
assert len(rows) == len(by_id) == len(decisions) == len(by_decision) == 342
assert set(by_id) == set(by_decision)
assert all(d["track"] in ("project", "commercial", "exclude") for d in decisions)

def eligible(req):
    return req["status"] in ("required", "task", "nice") and (
        req["category"] != "industry" or req["status"] == "required")

def key(track, req):
    return json.dumps([track, req["category"], req["concept"]], ensure_ascii=False, separators=(",", ":"))

def quoted_text(row, scope):
    if scope == "description":
        return row["description"] or ""
    assert scope.startswith("native:")
    return (row.get("native_sections") or {}).get(scope[7:], "") or ""

all_evidence = []
for decision in decisions:
    assert decision["reason"].strip()
    row = by_id[str(decision["id"])]
    bodies = [row["description"] or "", *(row.get("native_sections") or {}).values()]
    assert any(decision["role_evidence"] in body for body in bodies if isinstance(body, str)), (decision["id"], "Доказательство роли")
    if decision["track"] == "exclude":
        assert not decision["requirements"]
        continue
    for index, req in enumerate(decision["requirements"]):
        assert req["quote"] and req["quote"] in quoted_text(by_id[str(decision["id"])], req["scope"]), (
            decision["id"], index, "Цитата не найдена")
        assert req["status"] in ("required", "task", "nice", "unclear")
        all_evidence.append({
            "id": str(decision["id"]), "track": decision["track"],
            "index": index, "key": key(decision["track"], req), **req
        })

eligible_keys = {e["key"] for e in all_evidence if eligible(e)}
mapping = {}
group_ids = set()
for group in groups:
    assert group["id"] not in group_ids
    group_ids.add(group["id"])
    assert group["label"].strip() and group["track"] in ("project", "commercial")
    for member in group["member_keys"]:
        assert member not in mapping, ("Повторная группа", member)
        parsed = json.loads(member)
        assert parsed[0] == group["track"]
        mapping[member] = group["id"]
assert set(mapping) == eligible_keys, {
    "unmapped": list(eligible_keys - set(mapping)),
    "unknown": list(set(mapping) - eligible_keys)
}
evidence_by_group = {g["id"]: [] for g in groups}
for evidence in all_evidence:
    if eligible(evidence):
        evidence["group_id"] = mapping[evidence["key"]]
        evidence_by_group[evidence["group_id"]].append(evidence)
    else:
        evidence["group_id"] = None

counts = Counter(d["track"] for d in decisions)
assert dict(counts) == manifest["expected_tracks"]
active_counts = Counter(d["track"] for d in decisions if by_id[str(d["id"])]["is_active"])
summary = []
for group in groups:
    evidence = evidence_by_group[group["id"]]
    ids = {e["id"] for e in evidence}
    assert ids
    if group["category"] == "industry":
        assert all(e["category"] == "industry" and e["status"] == "required" for e in evidence)
    assert not any(e["category"] == "industry" for e in evidence) or group["category"] == "industry"
    nice_only = {e["id"] for e in evidence if e["status"] == "nice"} - {
        e["id"] for e in evidence if e["status"] in ("required", "task")}
    nice_any = {e["id"] for e in evidence if e["status"] == "nice"}
    examples = group["examples"]
    assert 1 <= len(examples) <= 3
    for example in examples:
        assert any(e["id"] == str(example["id"]) and e["quote"] == example["quote"]
                   and e["scope"] == example["scope"] for e in evidence), (group["id"], example)
    summary.append({
        **group, "count": len(ids), "total": counts[group["track"]],
        "active": sum(bool(by_id[i]["is_active"]) for i in ids),
        "nice": len(nice_any), "nice_only": len(nice_only)
    })

wb = Workbook()
wb.remove(wb.active)
wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
wb.properties.title = "Карьерный шкаф: проектный и коммерческий треки"
wb.properties.subject = "Карта требований из 342 сохранённых вакансий"
wb.properties.creator = "ProdRadar"
NAVY, GREEN, INK, LIGHT = "17324D", "236C57", "183042", "F1F5F8"
WHITE, ORANGE = "FFFFFF", "FFF0D5"
track_names = {"project": "Проектный трек", "commercial": "Коммерческий трек", "exclude": "Исключена"}
status_names = {"required": "Требование", "task": "Задача", "nice": "Будет плюсом", "unclear": "Неоднозначно"}
category_names = {
    "skill": "Навык", "tool": "Инструмент", "method": "Метод",
    "experience": "Опыт", "education": "Образование", "language": "Язык",
    "domain": "Предметная область", "industry": "Отраслевое требование"
}

def put(ws, row, column, value, fill=None, bold=False, color=INK, size=11):
    cell = ws.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Calibri", size=size, color=color, bold=bold)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell

def title(ws, text, color, last=10):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    put(ws, 1, 1, text, color, True, WHITE, 20)
    ws.row_dimensions[1].height = 36
    ws.sheet_view.showGridLines = False

def header(ws, row, values, color=NAVY):
    for col, value in enumerate(values, 1):
        put(ws, row, col, value, color, True, WHITE)
    ws.row_dimensions[row].height = 33
    ws.freeze_panes = f"C{row+1}"
    ws.auto_filter.ref = f"A{row}:{ws.cell(ws.max_row, len(values)).coordinate}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{row}"

overview = wb.create_sheet("Обзор")
title(overview, "Карьерный шкаф", NAVY, 5)
overview.merge_cells("A2:E2")
put(overview, 2, 1, "Два направления: управление проектами и коммерческое развитие", bold=True, size=13)
overview.row_dimensions[2].height = 25
overview.merge_cells("A3:E3")
put(overview, 3, 1, "Срез 14 октября 2026 · Все сохранённые вакансии тестового профиля, включая архив", size=12)
overview.row_dimensions[3].height = 32
header(overview, 5, ["Пул", "Вакансий", "Активны", "Компаний", "Что внутри"])
overview_rows = [
    ("Все сохранённые", len(rows), sum(r["is_active"] for r in rows),
     len({r["company"] for r in rows}), "Каждая вакансия разобрана один раз"),
    ("Проектный трек", counts["project"], active_counts["project"],
     len({by_id[str(d["id"])]["company"] for d in decisions if d["track"] == "project"}),
     "Планирование и ведение проектов, запусков и изменений"),
    ("Коммерческий трек", counts["commercial"], active_counts["commercial"],
     len({by_id[str(d["id"])]["company"] for d in decisions if d["track"] == "commercial"}),
     "Развитие бизнеса, клиентов и партнёрств, коммерческие процессы"),
    ("Исключены", counts["exclude"], active_counts["exclude"],
     len({by_id[str(d["id"])]["company"] for d in decisions if d["track"] == "exclude"}),
     "Продуктовые и другие роли вне двух направлений; причины — на листе «Вакансии»")
]
for row_index, values in enumerate(overview_rows, 6):
    for col, value in enumerate(values, 1):
        put(overview, row_index, col, value, LIGHT if row_index % 2 == 0 else None)
    overview.row_dimensions[row_index].height = 50
for row_index, message in enumerate([
    "Начните с нужного трека. Оцените каждый навык от 1 до 5, запишите примеры своего опыта и следующий шаг.",
    "Частота — доля вакансий своего трека. Повторы внутри одной вакансии не увеличивают счётчик.",
    "Общие пожелания учтены в частоте и видны в колонке «Из них — плюс». Отраслевые пожелания в частоту не входят.",
    "Цитаты взяты из исходных текстов. Для редких требований приведены доступные примеры.",
    "Лист «Доказательства» содержит все извлечённые формулировки и их статус. Лист «Вакансии» — весь исходный пул."
], 12):
    overview.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=5)
    put(overview, row_index, 1, message, size=12)
    overview.row_dimensions[row_index].height = 39
for col, width in {"A":32, "B":14, "C":14, "D":14, "E":64}.items():
    overview.column_dimensions[col].width = width
overview.auto_filter.ref = "A5:E9"
overview.freeze_panes = "B6"

for track, color in [("project", NAVY), ("commercial", GREEN)]:
    ws = wb.create_sheet(track_names[track])
    title(ws, track_names[track], color)
    ws.merge_cells("A2:J2")
    put(ws, 2, 1, f"Вакансий: {counts[track]} · Активны: {active_counts[track]} · Самооценка: 1 — нет опыта, 5 — уверенный опыт", bold=True, size=12)
    ws.row_dimensions[2].height = 27
    headers = ["Навык или требование", "Точные цитаты из вакансий", "Вакансий", "%",
               "Из них — плюс", "Самооценка 1–5", "Комментарий", "Действия", "Тип требования", "Активных"]
    for c, label in enumerate(headers, 1):
        put(ws, 4, c, label, color, True, WHITE)
    ws.row_dimensions[4].height = 34
    local = sorted([g for g in summary if g["track"] == track],
                   key=lambda g: (g["category"] == "industry", -g["count"], g["label"]))
    for r, group in enumerate(local, 5):
        examples = []
        for ex in group["examples"]:
            job = by_id[str(ex["id"])]
            examples.append(f'«{unescape(ex["quote"])}» — {job["company"]}')
        values = [group["label"], "\n\n".join(examples), group["count"],
                  f"=C{r}/$C$3", group["nice"], None, None, None,
                  category_names.get(group["category"], group["category"]), group["active"]]
        base = ORANGE if group["category"] == "industry" else (LIGHT if r % 2 else None)
        for c, value in enumerate(values, 1):
            put(ws, r, c, value, base if c not in (6, 7, 8) else "FFFBEF")
        ws.cell(r, 3).number_format = f'0" из {counts[track]}"'
        ws.cell(r, 4).data_type = "f"
        ws.cell(r, 4).number_format = "0%"
        ws.cell(r, 5).number_format = "0"
        ws.row_dimensions[r].height = min(409, max(80, 18 * sum(
            max(1, len(line) // 65 + 1) for line in values[1].split("\n"))))
    put(ws, 3, 1, "База для долей")
    put(ws, 3, 3, counts[track], bold=True)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:J{ws.max_row}"
    dv = DataValidation(type="whole", operator="between", formula1=1, formula2=5, allow_blank=True)
    dv.error = "Введите целое число от 1 до 5."
    dv.errorTitle = "Самооценка"
    dv.showErrorMessage = True
    dv.errorStyle = "stop"
    ws.add_data_validation(dv)
    dv.add(f"F5:F{ws.max_row}")
    ws.conditional_formatting.add(f"D5:D{ws.max_row}", ColorScaleRule(
        start_type="min", start_color="EDF3F6", end_type="max", end_color="A4C9C0"))
    for c, width in {"A":32,"B":66,"C":13,"D":8,"E":14,"F":14,"G":30,"H":33,"I":20,"J":12}.items():
        ws.column_dimensions[c].width = width
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

evidence_ws = wb.create_sheet("Доказательства")
title(evidence_ws, "Цитаты и исходные требования", NAVY, 9)
headers = ["Трек", "Группа", "Статус", "Цитата", "Компания", "Вакансия", "Ссылка", "Учтено в частоте", "Тип"]
for r, e in enumerate(all_evidence, 5):
    job = by_id[e["id"]]
    group = next((g for g in groups if g["id"] == e["group_id"]), None)
    values = [track_names[e["track"]], group["label"] if group else e["concept"],
              status_names[e["status"]], unescape(e["quote"]), job["company"], job["title"],
              job["url"], "Да" if eligible(e) else "Нет", category_names[e["category"]]]
    for c, value in enumerate(values, 1):
        put(evidence_ws, r, c, value, LIGHT if r % 2 else None)
    evidence_ws.cell(r, 7).hyperlink = job["url"]
    evidence_ws.cell(r, 7).font = Font(name="Calibri", size=11, color="165EA6", underline="single")
    evidence_ws.row_dimensions[r].height = min(409, max(44, len(e["quote"]) // 90 * 16 + 44))
header(evidence_ws, 4, headers)
for c, width in {"A":22,"B":38,"C":20,"D":92,"E":20,"F":55,"G":36,"H":20,"I":24}.items():
    evidence_ws.column_dimensions[c].width = width

jobs_ws = wb.create_sheet("Вакансии")
title(jobs_ws, "Все вакансии и решения по ним", NAVY, 8)
exclusion_names = {"product":"Продуктовая роль","other":"Другая роль","insufficient":"Недостаточно текста"}
for r, job in enumerate(rows, 5):
    d = by_decision[str(job["id"])]
    values = [job["company"], job["title"], track_names[d["track"]],
              "Да" if job["is_active"] else "Нет", d["reason"], job["url"],
              exclusion_names.get(d.get("exclusion"), ""), str(job["id"])]
    for c, value in enumerate(values, 1):
        put(jobs_ws, r, c, value, LIGHT if r % 2 else None)
    jobs_ws.cell(r, 6).hyperlink = job["url"]
    jobs_ws.cell(r, 6).font = Font(name="Calibri", size=11, color="165EA6", underline="single")
    jobs_ws.row_dimensions[r].height = 66
header(jobs_ws, 4, ["Компания","Вакансия","Трек","Активна","Почему","Ссылка","Причина исключения","ID"])
for c, width in {"A":22,"B":60,"C":23,"D":12,"E":85,"F":38,"G":26,"H":45}.items():
    jobs_ws.column_dimensions[c].width = width

out = Path("output")
out.mkdir(exist_ok=True)
filename = out / "Карьерный_шкаф.xlsx"
wb.save(filename)
checked = load_workbook(filename, data_only=False)
assert checked.sheetnames == ["Обзор","Проектный трек","Коммерческий трек","Доказательства","Вакансии"]
assert checked["Вакансии"].max_row - 4 == 342
assert checked["Доказательства"].max_row - 4 == len(all_evidence)
for track in ("project", "commercial"):
    ws = checked[track_names[track]]
    assert ws.cell(3,3).value == counts[track]
    assert ws.max_row - 4 == sum(g["track"] == track for g in summary)
    for r in range(5, ws.max_row + 1):
        assert ws.cell(r, 4).value == f"=C{r}/$C$3"
        assert all(ws.cell(r,c).value is None for c in (6,7,8))
        assert 0 <= ws.cell(r,5).value <= ws.cell(r,3).value <= counts[track]
assert counts["project"] + counts["commercial"] + counts["exclude"] == 342
report = {
    "total":342,"active":sum(r["is_active"] for r in rows),
    "tracks":dict(counts),"active_tracks":dict(active_counts),
    "groups":Counter(g["track"] for g in groups),
    "evidence":len(all_evidence),"industry_groups":sum(g["category"]=="industry" for g in groups),
    "industry_nice_not_counted":sum(e["category"]=="industry" and e["status"]=="nice" for e in all_evidence),
    "quotes_verified":len(all_evidence),"filename":filename.name,
    "source_sha256":source["meta"]["source_sha256"],
    "checks":"Все ID, цитаты, группы, знаменатели, формулы и пустые поля самооценки проверены"
}
(out / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print("CAREER_FINAL=" + json.dumps(report, ensure_ascii=False))
